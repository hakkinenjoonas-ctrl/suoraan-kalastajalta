from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from reportlab.graphics.barcode import qr
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx"
OUT.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT / "MUNBYN_RW403B_pikaohje_kalastajalle.docx"
QR_PATH = ROOT / "tmp" / "munbyn-print-google-play-qr.png"
PLAY_STORE_URL = "https://play.google.com/store/apps/details?id=com.syz.mprint"
LOGO_DIR = Path("/Users/joonashakkinen/Desktop/Suoraan Kalastajalta APP/Logot")
LOGO_MATCHES = sorted(LOGO_DIR.glob("Logo tekstill*.png"))
LOGO_PATH = LOGO_MATCHES[0] if LOGO_MATCHES else ROOT / "public" / "logo.png"

BLUE = "174B74"
BLUE_DARK = "0F3554"
PALE = "EAF4FA"
PALE_GREEN = "EAF7F1"
GREEN = "167A58"
ORANGE = "E88032"
TEXT = "17324A"
MUTED = "51687A"
WHITE = "FFFFFF"


def create_qr_code():
    QR_PATH.parent.mkdir(parents=True, exist_ok=True)
    widget = qr.QrCodeWidget(PLAY_STORE_URL)
    widget.getBounds()
    modules = widget.qr.modules
    quiet_zone = 4
    scale = 10
    module_count = len(modules)
    size = (module_count + quiet_zone * 2) * scale
    image = Image.new("RGB", (size, size), "white")
    draw = ImageDraw.Draw(image)
    for row_index, row in enumerate(modules):
        for column_index, filled in enumerate(row):
            if filled:
                x = (column_index + quiet_zone) * scale
                y = (row_index + quiet_zone) * scale
                draw.rectangle((x, y, x + scale - 1, y + scale - 1), fill="black")
    image.save(QR_PATH, dpi=(300, 300))


create_qr_code()


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, size=9.0, bold=False, color=TEXT, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_step(cell, number, title, body, note=None):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    badge = p.add_run(f"{number}  ")
    badge.bold = True
    badge.font.name = "Aptos Display"
    badge.font.size = Pt(13)
    badge.font.color.rgb = RGBColor.from_string(ORANGE)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(11.2)
    r.font.color.rgb = RGBColor.from_string(BLUE_DARK)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(4 if note else 7)
    p2.paragraph_format.left_indent = Mm(5)
    p2.paragraph_format.line_spacing = 1.04
    r2 = p2.add_run(body)
    r2.font.name = "Aptos"
    r2.font.size = Pt(9.4)
    r2.font.color.rgb = RGBColor.from_string(TEXT)
    if note:
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(7)
        p3.paragraph_format.left_indent = Mm(5)
        rr = p3.add_run(note)
        rr.italic = True
        rr.font.name = "Aptos"
        rr.font.size = Pt(8.5)
        rr.font.color.rgb = RGBColor.from_string(MUTED)


doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = Mm(10)
sec.bottom_margin = Mm(9)
sec.left_margin = Mm(12)
sec.right_margin = Mm(12)
sec.header_distance = Mm(4)
sec.footer_distance = Mm(4)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9)
styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)
styles["Normal"].paragraph_format.space_after = Pt(3)

# Header
header = doc.add_table(rows=1, cols=2)
header.autofit = False
header.columns[0].width = Mm(124)
header.columns[1].width = Mm(62)
no_table_borders(header)
left, right = header.rows[0].cells
cell_margins(left, 0, 0, 40, 80)
cell_margins(right, 0, 0, 20, 0)
p = left.paragraphs[0]
p.paragraph_format.space_after = Pt(1)
r = p.add_run("MUNBYN RW403B")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(10)
r.font.color.rgb = RGBColor.from_string(ORANGE)
p = left.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Etikettitulostimen pikaohje")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(23)
r.font.color.rgb = RGBColor.from_string(BLUE_DARK)
p = left.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Yhdistä tulostin puhelimeen ja tulosta kalaerän etiketti.")
r.font.name = "Aptos"
r.font.size = Pt(10)
r.font.color.rgb = RGBColor.from_string(MUTED)
if LOGO_PATH.exists():
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    run.add_picture(str(LOGO_PATH), width=Mm(55))

# Important warning
warning = doc.add_table(rows=1, cols=1)
warning.autofit = False
warning.columns[0].width = Mm(186)
no_table_borders(warning)
cell = warning.cell(0, 0)
shade(cell, PALE_GREEN)
cell_margins(cell, 100, 180, 100, 180)
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
r = p.add_run("TÄRKEÄÄ  ")
r.bold = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor.from_string(GREEN)
r = p.add_run("Älä yhdistä tulostinta puhelimen Bluetooth-valikossa. Yhteys tehdään Munbyn Print -sovelluksen sisällä.")
r.bold = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor.from_string(TEXT)

doc.add_paragraph().paragraph_format.space_after = Pt(0)

# Two-column body
body = doc.add_table(rows=1, cols=2)
body.autofit = False
body.columns[0].width = Mm(91)
body.columns[1].width = Mm(91)
no_table_borders(body)
c1, c2 = body.rows[0].cells
for c in (c1, c2):
    cell_margins(c, 70, 140, 70, 140)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

shade(c1, PALE)
set_cell_text(c1, "A. OTA TULOSTIN KÄYTTÖÖN", 10.2, True, BLUE)
add_step(c1, "1", "Laita etikettirulla paikalleen", "Avaa kansi ja aseta etikettirulla tulostimeen niin, että etiketit tulevat ulos tulostusaukosta. Säädä paperiohjaimet etiketin reunoihin ja sulje kansi. Katso tarvittaessa asennuskuvat tulostimen mukana toimitetusta MUNBYN-ohjeesta.")
add_step(c1, "2", "Kytke virta", "Liitä virtajohto tulostimeen ja pistorasiaan. Käynnistä tulostin. Pidä tulostin lähellä puhelinta.")
add_step(c1, "3", "Lataa Munbyn Print", "Avaa puhelimen Google Play -kauppa (Android) tai App Store (iPhone). Hae “Munbyn Print” ja asenna sovellus.")
add_step(c1, "4", "Salli Bluetooth", "Laita puhelimen Bluetooth päälle. Salli Munbyn Printille Bluetooth- ja lähilaitteiden käyttö, jos puhelin kysyy lupaa.")
add_step(c1, "5", "Yhdistä RW403B", "Avaa Munbyn Print. Valitse sovelluksessa laite/tulostin, etsi “RW403B” ja napauta sitä. Hyväksy yhteys valitsemalla OK.", "Tulostimen tila näkyy sovelluksessa yhdistettynä (Connected).")
p = c1.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(1)
p.paragraph_format.space_after = Pt(1)
r = p.add_run()
r.add_picture(str(QR_PATH), width=Mm(19))
p = c1.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Skannaa: Munbyn Print / Google Play")
r.bold = True
r.font.size = Pt(7.8)
r.font.color.rgb = RGBColor.from_string(BLUE)

shade(c2, "F5F8FA")
set_cell_text(c2, "B. TULOSTA ETIKETTI SOVELLUKSESTA", 10.2, True, BLUE)
add_step(c2, "1", "Avaa tallennettu kalaerä", "Avaa Suoraan Kalastajalta -sovelluksessa Saaliit-näkymä ja valitse kalaerä, jolle haluat etiketin.")
add_step(c2, "2", "Valitse Tulosta etiketit", "Napauta kalaerän kohdalla “Tulosta etiketit”.")
add_step(c2, "3", "Tarkista etiketin tiedot", "Valitse vesityyppi (Makea vesi tai Meri), laatikoiden määrä ja tarvittaessa paino, tuotemuoto sekä viimeinen käyttöpäivä.")
add_step(c2, "4", "Valitse oikea pohja", "Valitse tulostuspohjaksi “MUNBYN 4x3” (101,6 × 76,2 mm). Tarkista etiketti esikatselusta.")
add_step(c2, "5", "Avaa etiketti Munbyn Printissä", "Napauta “Luo PDF” tai “Tulosta” - molemmat avaavat puhelimen näytölle saman sovellusvalikon. Valitse valikosta Munbyn Print. Jos sovellus ei näy heti, valitse Jaa tai Avaa sovelluksessa ja sen jälkeen Munbyn Print.")
add_step(c2, "6", "Tarkista asetukset ja tulosta", "Varmista Munbyn Printissä, että RW403B on yhdistetty, paperikoko on 4 × 3 tuumaa (101,6 × 76,2 mm), suunta on oikein ja kopiomäärä vastaa laatikoiden määrää. Napauta Print / Tulosta.")

doc.add_paragraph().paragraph_format.space_after = Pt(0)

# Troubleshooting strip
tr = doc.add_table(rows=1, cols=3)
tr.autofit = False
tr.columns[0].width = Mm(62)
tr.columns[1].width = Mm(62)
tr.columns[2].width = Mm(62)
no_table_borders(tr)
tips = [
    ("Tulostin ei löydy", "Tarkista virta ja Bluetooth. Sulje Munbyn Print, avaa se uudelleen ja etsi RW403B sovelluksessa."),
    ("Etiketti tulostuu väärin", "Tarkista MUNBYN 4x3, paperikoko 101,6 × 76,2 mm ja ettei tulostuksessa ole sovitusta tai reunuksia."),
    ("Paperi ei osu kohdalleen", "Aseta rulla suoraksi ja paperiohjaimet etiketin reunoihin. Sulje kansi napakasti."),
]
for cell, (title, body_text) in zip(tr.rows[0].cells, tips):
    shade(cell, BLUE_DARK)
    cell_margins(cell, 90, 120, 90, 120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(body_text)
    r.font.size = Pt(8.4)
    r.font.color.rgb = RGBColor.from_string(WHITE)

contact = doc.add_table(rows=1, cols=1)
contact.autofit = False
contact.columns[0].width = Mm(186)
no_table_borders(contact)
cell = contact.cell(0, 0)
shade(cell, PALE_GREEN)
cell_margins(cell, 90, 150, 90, 150)
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
r = p.add_run("KYSY APUA TARVITTAESSA")
r.bold = True
r.font.size = Pt(8.2)
r.font.color.rgb = RGBColor.from_string(GREEN)
p = cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Joonas Häkkinen  •  info@suoraankalastajalta.fi  •  040 735 4021")
r.bold = True
r.font.size = Pt(9.2)
r.font.color.rgb = RGBColor.from_string(BLUE_DARK)

footer = doc.add_paragraph()
footer.paragraph_format.space_before = Pt(4)
footer.paragraph_format.space_after = Pt(0)
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Suoraan Kalastajalta  •  MUNBYN RW403B  •  Pikaohje kalastajalle")
r.font.size = Pt(7.5)
r.font.color.rgb = RGBColor.from_string(MUTED)

doc.save(DOCX_PATH)
print(DOCX_PATH)
